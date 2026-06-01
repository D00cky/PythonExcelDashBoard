"""Walk a .docx template and dump every placeholder-like text run + image anchor.

Run:  .venv/bin/python scripts/inspect_docx_template.py <path-to-docx> [more...]
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.document import Document as _Doc
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

PLACEHOLDER_RE = re.compile(r"\{\{[^}]+\}\}|\[\[[^\]]+\]\]|<<[^>]+>>")


def _iter_paragraphs(doc: _Doc):
    for p in doc.paragraphs:
        yield ("body", p)
    for ti, table in enumerate(doc.tables):
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                for p in cell.paragraphs:
                    yield (f"table[{ti}].cell[{ri},{ci}]", p)


def _runs_text(p: Paragraph) -> list[str]:
    return [r.text for r in p.runs]


def _has_sdt(p: Paragraph) -> bool:
    return p._element.find(qn("w:sdt")) is not None or any(
        child.tag == qn("w:sdt") for child in p._element.iter()
    )


def _embedded_images(doc: _Doc) -> list[str]:
    images = []
    for rel_id, rel in doc.part.rels.items():
        if "image" in rel.reltype:
            images.append(f"{rel_id}: {rel.target_ref}")
    return images


def inspect(path: Path) -> None:
    print(f"\n{'=' * 80}\nFILE: {path.name}\n{'=' * 80}")
    doc = Document(str(path))

    placeholders: dict[str, list[str]] = {}
    sdt_locations: list[str] = []
    all_paragraphs = 0
    nonempty_paragraphs = 0

    for loc, p in _iter_paragraphs(doc):
        all_paragraphs += 1
        full = p.text
        if not full.strip():
            continue
        nonempty_paragraphs += 1
        for m in PLACEHOLDER_RE.findall(full):
            placeholders.setdefault(m, []).append(loc)
        if _has_sdt(p):
            sdt_locations.append(loc)

    print(f"Total paragraphs (body+tables): {all_paragraphs} (non-empty: {nonempty_paragraphs})")
    print(f"Tables: {len(doc.tables)}")
    print(f"Embedded images: {len(_embedded_images(doc))}")
    for img in _embedded_images(doc)[:20]:
        print(f"  IMG  {img}")

    print(f"\nPlaceholders found: {len(placeholders)}")
    for ph, locs in sorted(placeholders.items()):
        print(f"  {ph!r}  ×{len(locs)}  e.g. {locs[0]}")

    print(f"\nContent-control (w:sdt) paragraphs: {len(sdt_locations)}")
    for loc in sdt_locations[:10]:
        print(f"  SDT  {loc}")

    print("\n--- First 80 non-empty paragraphs (preview) ---")
    shown = 0
    for loc, p in _iter_paragraphs(doc):
        t = p.text.strip()
        if not t:
            continue
        runs = _runs_text(p)
        run_repr = " | ".join(repr(r) for r in runs if r)
        print(f"[{loc}] text={t!r}")
        if len(runs) > 1:
            print(f"   runs={run_repr}")
        shown += 1
        if shown >= 80:
            break


def main() -> int:
    if len(sys.argv) < 2:
        print(__doc__)
        return 1
    for arg in sys.argv[1:]:
        inspect(Path(arg))
    return 0


if __name__ == "__main__":
    sys.exit(main())
