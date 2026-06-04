from io import BytesIO

import pytest
from docx import Document
from openpyxl import load_workbook

from app.core.aggregator import PoloBatch, PoloFile
from app.core.exporters.docx_sabesp import (
    MENSAL_SKELETON_PATH,
    SEMANAL_SKELETON_PATH,
    IndiceRow,
    MensalContext,
    SemanalContext,
    batch_context_from_batch,
    context_from_template,
    indice_rows_from_inspections,
    render_mensal,
    render_semanal,
    semanal_context_from_template,
)
from app.core.templates.pimentas import PimentasTemplate
from tests.fixtures.pimentas_minimal import make_minimal_pimentas


def _all_text(doc):
    chunks = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    chunks.append(p.text)
    return "\n".join(chunks)


def test_render_mensal_substitutes_polo_and_period():
    ctx = MensalContext(
        polo_label="Pimentas",
        periodo_inicio="05-06-2026",
        periodo_fim="05-12-2026",
        mes_extenso="Maio",
        ano="2026",
    )

    body = render_mensal(ctx)

    assert body[:2] == b"PK"  # .docx is a zip
    doc = Document(BytesIO(body))
    text = _all_text(doc)
    assert "Pimentas" in text
    assert "PIMENTAS" in text
    assert "05-06-2026" in text
    assert "05-12-2026" in text
    assert "Maio de 2026" in text
    # No source polo name should leak through.
    assert "Gopoúva" not in text
    assert "Gopouva" not in text
    assert "GOPOÚVA" not in text
    # No unrendered Jinja markers.
    assert "{{" not in text


def test_render_mensal_polo_label_upper_derived_from_polo_label():
    ctx = MensalContext(
        polo_label="Santana",
        periodo_inicio="06-01-2026",
        periodo_fim="06-07-2026",
        mes_extenso="Junho",
        ano="2026",
    )
    assert ctx.polo_label_upper == "SANTANA"


def test_context_from_template_uses_inspection_date_bounds(tmp_path):
    path = make_minimal_pimentas(tmp_path, with_inspections=True)
    wb = load_workbook(path, data_only=True, read_only=True)
    template = PimentasTemplate.detect(wb.sheetnames)

    ctx = context_from_template(template, path)

    # Fixture dates span 2026-03-05 to 2026-03-29 — period text and month should reflect that.
    assert ctx.periodo_inicio == "03-05-2026"
    assert ctx.periodo_fim == "03-29-2026"
    assert ctx.mes_extenso == "Março"
    assert ctx.ano == "2026"
    assert ctx.polo_label == "Pimentas"


def test_context_from_template_handles_workbook_without_inspections(tmp_path):
    path = make_minimal_pimentas(tmp_path, with_inspections=False)
    wb = load_workbook(path, data_only=True, read_only=True)
    template = PimentasTemplate.detect(wb.sheetnames)

    ctx = context_from_template(template, path)

    # No dated rows → blank period, but polo label still binds.
    assert ctx.polo_label == "Pimentas"
    assert ctx.periodo_inicio == ""
    assert ctx.periodo_fim == ""
    assert ctx.mes_extenso == ""


@pytest.mark.parametrize(
    "field",
    ["polo_label", "periodo_inicio", "periodo_fim", "mes_extenso", "ano", "indice"],
)
def test_render_dict_exposes_every_skeleton_placeholder(field):
    ctx = MensalContext(
        polo_label="X",
        periodo_inicio="i",
        periodo_fim="f",
        mes_extenso="m",
        ano="a",
    )
    assert field in ctx.as_render_dict()
    assert "polo_label_upper" in ctx.as_render_dict()


def _indice_table(doc):
    return next(t for t in doc.tables if "ÍNDICE TECNOLÓGICO" in t.rows[0].cells[0].text)


def test_render_mensal_strips_manual_audit_sections():
    """The digital-surveillance scope drops §2 DESCRIÇÃO DA EQUIPE,
    §3 PLANO DE AMOSTRAGEM, §4 AUDITORIAS REALIZADAS / DCP, §5 MAPEAMENTO,
    §6.1-6.3 (manual-audit narratives), §6.5-6.6 (NC do canteiro photos),
    §7 NÃO CONFORMIDADES. The renderer must produce a doc that has none of
    those headings nor their characteristic content."""
    ctx = MensalContext(
        polo_label="Pimentas",
        periodo_inicio="05-06-2026",
        periodo_fim="05-12-2026",
        mes_extenso="Maio",
        ano="2026",
    )

    doc = Document(BytesIO(render_mensal(ctx)))
    text = _all_text(doc)

    for stripped in (
        "DESCRIÇÃO DA EQUIPE",
        "PLANO DE AMOSTRAGEM",
        "AUDITORIAS REALIZADAS",
        "5. MAPEAMENTO",  # §5 only; §8.1 also says MAPEAMENTO so we anchor on prefix
        "6.1 QUANTIDADE",
        "6.2 ÍNDICE DE CONTROLE",
        "6.5 NÃO CONFORMIDADES",
        "EQUIPE I – Tecnólogo",  # original Sondotécnica auditor org chart
        "CONTROLE DE DCP",
        "Lucas Jeremias",  # Sondotécnica auditor name from the source
    ):
        assert stripped not in text, f"{stripped!r} should have been stripped"


def test_render_mensal_keeps_digital_surveillance_sections():
    """The kept sections are the cover, §1 INTRODUÇÃO, §6.4 (ÍNDICE TECNOLÓGICO
    POR EQUIPE table lives here), §8 ACOMPANHAMENTO – OLHAR DIGITAL plus 8.x
    subsections, and §9 CONCLUSÃO."""
    ctx = MensalContext(
        polo_label="Pimentas",
        periodo_inicio="05-06-2026",
        periodo_fim="05-12-2026",
        mes_extenso="Maio",
        ano="2026",
    )

    doc = Document(BytesIO(render_mensal(ctx)))
    text = _all_text(doc)

    for kept in (
        "1. INTRODUÇÃO",
        "6.4 ÍNDICE DE CONFORMIDADE POR EQUIPE",
        "8. ACOMPANHAMENTO",
        "OLHAR DIGITAL",
        "8.1 MAPEAMENTO FISCALIZAÇÕES",
        "8.3 GRÁFICO IQS",
        "9. CONCLUSÃO",
    ):
        assert kept in text, f"{kept!r} should still be present"


def test_render_mensal_indice_table_repeats_one_row_per_supplied_entry():
    ctx = MensalContext(
        polo_label="Pimentas",
        periodo_inicio="05-06-2026",
        periodo_fim="05-12-2026",
        mes_extenso="Maio",
        ano="2026",
        indice_tecnologico=(
            IndiceRow(equipe="Maria", servico="Água", quantidade=12, ic_pct_str="98,50"),
            IndiceRow(equipe="João", servico="Esgoto", quantidade=5, ic_pct_str="100,00"),
            IndiceRow(equipe="Ana", servico="Reposição", quantidade=8, ic_pct_str="95,25"),
        ),
    )

    doc = Document(BytesIO(render_mensal(ctx)))
    table = _indice_table(doc)
    # 3 header rows + 3 data rows.
    assert len(table.rows) == 6
    data_rows = [[c.text.strip() for c in row.cells] for row in table.rows[3:]]
    assert data_rows == [
        ["Maria", "Água", "12", "98,50"],
        ["João", "Esgoto", "5", "100,00"],
        ["Ana", "Reposição", "8", "95,25"],
    ]


def test_render_mensal_indice_table_collapses_to_zero_rows_when_empty():
    ctx = MensalContext(
        polo_label="Pimentas",
        periodo_inicio="05-06-2026",
        periodo_fim="05-12-2026",
        mes_extenso="Maio",
        ano="2026",
        indice_tecnologico=(),
    )

    doc = Document(BytesIO(render_mensal(ctx)))
    table = _indice_table(doc)
    # Just the 3 header rows survive.
    assert len(table.rows) == 3


def test_indice_rows_from_inspections_groups_by_team_and_service(tmp_path):
    path = make_minimal_pimentas(tmp_path, with_inspections=True)
    wb = load_workbook(path, data_only=True, read_only=True)
    template = PimentasTemplate.detect(wb.sheetnames)

    rows = indice_rows_from_inspections(template.extract_inspections(path))

    assert rows  # non-empty for the fixture
    # Every row has positive quantidade and a pt-BR formatted IC%.
    for r in rows:
        assert r.quantidade > 0
        assert "," in r.ic_pct_str
        # IC% strictly in [0.00, 100.00] formatted with two decimals.
        whole, frac = r.ic_pct_str.split(",")
        assert 0 <= int(whole) <= 100
        assert len(frac) == 2


def test_indice_rows_from_inspections_returns_empty_for_workbook_without_inspections(tmp_path):
    path = make_minimal_pimentas(tmp_path, with_inspections=False)
    wb = load_workbook(path, data_only=True, read_only=True)
    template = PimentasTemplate.detect(wb.sheetnames)

    rows = indice_rows_from_inspections(template.extract_inspections(path))
    assert rows == ()


def test_context_from_template_populates_indice_tecnologico_from_xlsx(tmp_path):
    path = make_minimal_pimentas(tmp_path, with_inspections=True)
    wb = load_workbook(path, data_only=True, read_only=True)
    template = PimentasTemplate.detect(wb.sheetnames)

    ctx = context_from_template(template, path)

    assert ctx.indice_tecnologico  # populated automatically from xlsx


def test_context_from_template_populates_chart_pngs_for_section_8_1(tmp_path):
    path = make_minimal_pimentas(tmp_path, with_inspections=True)
    wb = load_workbook(path, data_only=True, read_only=True)
    template = PimentasTemplate.detect(wb.sheetnames)

    ctx = context_from_template(template, path)

    # Every §8.1 placeholder gets PNG bytes; kaleido emits at least a few KB.
    assert set(ctx.chart_pngs) == {"ic_bar", "iqs_bar", "photo_conformity"}
    for key, png in ctx.chart_pngs.items():
        assert png[:8] == b"\x89PNG\r\n\x1a\n", f"{key} should be a PNG"
        assert len(png) > 1024


def test_context_from_template_returns_empty_chart_pngs_on_kaleido_failure(tmp_path, monkeypatch):
    """Kaleido can OOM or crash on large workbooks (~4.6k inspections in the
    bug report). The exporter must degrade to empty PNG bytes — render_mensal
    skips empty pngs — instead of letting the exception bubble out and produce
    a blank HTTP response.
    """
    import plotly.graph_objects as go

    def boom(self, *a, **kw):
        raise RuntimeError("kaleido crashed")

    monkeypatch.setattr(go.Figure, "to_image", boom)

    path = make_minimal_pimentas(tmp_path, with_inspections=True)
    wb = load_workbook(path, data_only=True, read_only=True)
    template = PimentasTemplate.detect(wb.sheetnames)

    ctx = context_from_template(template, path)

    assert set(ctx.chart_pngs) == {"ic_bar", "iqs_bar", "photo_conformity"}
    for key, png in ctx.chart_pngs.items():
        assert png == b"", f"{key} should fall back to empty bytes, got {len(png)}"


def test_render_mensal_skips_empty_chart_pngs_without_error():
    """Regression guard for the `if png:` branch in render_mensal: empty bytes
    for a chart placeholder must render cleanly with no InlineImage added and
    no Jinja-undefined error."""
    ctx = MensalContext(
        polo_label="Pimentas",
        periodo_inicio="05-06-2026",
        periodo_fim="05-12-2026",
        mes_extenso="Maio",
        ano="2026",
        chart_pngs={"ic_bar": b"", "iqs_bar": b"", "photo_conformity": b""},
    )

    skeleton = Document(str(MENSAL_SKELETON_PATH))
    baseline = sum("image" in r.reltype for r in skeleton.part.rels.values())

    body = render_mensal(ctx)
    assert body[:2] == b"PK"
    doc = Document(BytesIO(body))
    rendered = sum("image" in r.reltype for r in doc.part.rels.values())
    # No new image relationships when every chart png is empty.
    assert rendered == baseline


def test_render_mensal_embeds_chart_images_under_section_8_1(tmp_path):
    path = make_minimal_pimentas(tmp_path, with_inspections=True)
    wb = load_workbook(path, data_only=True, read_only=True)
    template = PimentasTemplate.detect(wb.sheetnames)

    ctx = context_from_template(template, path)
    body = render_mensal(ctx)

    doc = Document(BytesIO(body))
    # Embedded images surface as docx image relationships.
    image_rels = [r for r in doc.part.rels.values() if "image" in r.reltype]
    # Three §8.1 KPI charts → at least three image relationships.
    assert len(image_rels) >= 3
    # And no unrendered chart placeholders.
    text = _all_text(doc)
    for key in ("ic_bar", "iqs_bar", "photo_conformity"):
        assert "{{ " + key not in text


# Semanal tests — cycle 2 (cover-page-only skeleton, polo + period bindings).


def test_render_semanal_substitutes_polo_and_period_in_cover_table():
    ctx = SemanalContext(
        polo_label="Pimentas",
        periodo_inicio="05-06-2026",
        periodo_fim="05-12-2026",
    )

    body = render_semanal(ctx)

    assert body[:2] == b"PK"  # docx zip
    doc = Document(BytesIO(body))
    # The cover table is doc.tables[0]; row 20 cell 1 is the period line,
    # row 34 cell 1 is the POLO line. Other merged cells hold duplicates.
    text = _all_text(doc)
    assert "05-06-2026" in text
    assert "05-12-2026" in text
    assert "PIMENTAS" in text
    # Source-template values must not leak.
    assert "03-01-2026" not in text
    assert "03-25-2026" not in text
    assert "EXTREMO NORTE" not in text
    # No unrendered Jinja markers.
    assert "{{" not in text


def test_semanal_context_polo_label_upper_derived_from_polo_label():
    ctx = SemanalContext(
        polo_label="Santana",
        periodo_inicio="05-04-2026",
        periodo_fim="05-10-2026",
    )
    assert ctx.polo_label_upper == "SANTANA"
    assert "polo_label_upper" in ctx.as_render_dict()


def test_semanal_context_from_template_uses_inspection_date_bounds(tmp_path):
    path = make_minimal_pimentas(tmp_path, with_inspections=True)
    wb = load_workbook(path, data_only=True, read_only=True)
    template = PimentasTemplate.detect(wb.sheetnames)

    ctx = semanal_context_from_template(template, path)

    # Fixture inspections span 2026-03-05 to 2026-03-29.
    assert ctx.periodo_inicio == "03-05-2026"
    assert ctx.periodo_fim == "03-29-2026"
    assert ctx.polo_label == "Pimentas"


def test_semanal_context_from_template_handles_workbook_without_inspections(tmp_path):
    path = make_minimal_pimentas(tmp_path, with_inspections=False)
    wb = load_workbook(path, data_only=True, read_only=True)
    template = PimentasTemplate.detect(wb.sheetnames)

    ctx = semanal_context_from_template(template, path)

    # No dated inspections → blank period strings, polo label still binds.
    assert ctx.polo_label == "Pimentas"
    assert ctx.periodo_inicio == ""
    assert ctx.periodo_fim == ""


def test_semanal_context_from_template_populates_chart_pngs(tmp_path):
    path = make_minimal_pimentas(tmp_path, with_inspections=True)
    wb = load_workbook(path, data_only=True, read_only=True)
    template = PimentasTemplate.detect(wb.sheetnames)

    ctx = semanal_context_from_template(template, path)

    # Same three dashboard charts as Mensal — ic_bar under §1.1.1, iqs_bar +
    # photo_conformity under §2.
    assert set(ctx.chart_pngs) == {"ic_bar", "iqs_bar", "photo_conformity"}
    for key, png in ctx.chart_pngs.items():
        assert png[:8] == b"\x89PNG\r\n\x1a\n", f"{key} should be a PNG"
        assert len(png) > 1024


def _batch_with_two_polos(tmp_path) -> PoloBatch:
    """Build an in-memory PoloBatch with two minimal Pimentas-shaped fixtures
    sharing the same fixture period but tagged with different polo names. The
    Sabesp Mensal aggregator only cares about file_path + polo on each PoloFile,
    so iso_week/month/period_start/end use the fixture's own dates."""
    from datetime import date

    pim = make_minimal_pimentas(
        tmp_path, polo="PIMENTAS", with_inspections=True, file_name="pim.xlsx"
    )
    san = make_minimal_pimentas(
        tmp_path, polo="SANTANA", with_inspections=True, file_name="san.xlsx"
    )
    files = [
        PoloFile(
            file_path=pim,
            polo="pimentas",
            iso_week="2026-W10",
            month="2026-03",
            period_start=date(2026, 3, 5),
            period_end=date(2026, 3, 29),
        ),
        PoloFile(
            file_path=san,
            polo="santana",
            iso_week="2026-W10",
            month="2026-03",
            period_start=date(2026, 3, 5),
            period_end=date(2026, 3, 29),
        ),
    ]
    return PoloBatch(batch_dir=tmp_path, files=files)


def test_batch_context_from_batch_uses_combined_inspections_period(tmp_path):
    batch = _batch_with_two_polos(tmp_path)

    ctx = batch_context_from_batch(batch)

    # Period derived from combined inspections (fixture spans 03-05 → 03-29).
    assert ctx.periodo_inicio == "03-05-2026"
    assert ctx.periodo_fim == "03-29-2026"
    assert ctx.mes_extenso == "Março"
    assert ctx.ano == "2026"


def test_batch_context_polo_label_lists_all_polos_when_multiple(tmp_path):
    batch = _batch_with_two_polos(tmp_path)

    ctx = batch_context_from_batch(batch)

    # Multi-polo batch surfaces every polo in the label so the cover page shows
    # the full scope at a glance instead of the legacy "Múltiplos Polos".
    assert "Pimentas" in ctx.polo_label
    assert "Santana" in ctx.polo_label
    assert ctx.polo_label.startswith("Polos:")


def test_batch_context_polo_label_uses_single_name_for_single_polo_batch(tmp_path):
    from datetime import date

    pim = make_minimal_pimentas(
        tmp_path, polo="PIMENTAS", with_inspections=True, file_name="solo.xlsx"
    )
    batch = PoloBatch(
        batch_dir=tmp_path,
        files=[
            PoloFile(
                file_path=pim,
                polo="pimentas",
                iso_week="2026-W10",
                month="2026-03",
                period_start=date(2026, 3, 5),
                period_end=date(2026, 3, 29),
            )
        ],
    )

    ctx = batch_context_from_batch(batch)

    assert ctx.polo_label == "Pimentas"


def test_batch_context_populates_chart_pngs_for_aggregate_summary(tmp_path):
    batch = _batch_with_two_polos(tmp_path)

    ctx = batch_context_from_batch(batch)

    # Same three §8.1 charts as the single-Polo Mensal, fed by combined data.
    assert set(ctx.chart_pngs) == {"ic_bar", "iqs_bar", "photo_conformity"}
    for key, png in ctx.chart_pngs.items():
        assert png[:8] == b"\x89PNG\r\n\x1a\n", f"{key} should be a PNG"


def test_batch_context_indice_rows_prefix_team_with_polo(tmp_path):
    batch = _batch_with_two_polos(tmp_path)

    ctx = batch_context_from_batch(batch)

    # Two polos share the same fixture team names; the aggregate indice must
    # disambiguate by prefixing the polo so reviewers can tell rows apart.
    equipes = [r.equipe for r in ctx.indice_tecnologico]
    assert equipes, "indice should be populated"
    pimentas_rows = [e for e in equipes if e.startswith("Pimentas")]
    santana_rows = [e for e in equipes if e.startswith("Santana")]
    assert pimentas_rows, f"expected Pimentas-prefixed rows, got {equipes}"
    assert santana_rows, f"expected Santana-prefixed rows, got {equipes}"


def test_indice_rows_from_inspections_group_by_polo_disambiguates_teams():
    """Direct unit test of the group_by_polo flag — bypasses the batch
    plumbing so the prefixing rule itself stays pinned."""
    import pandas as pd

    df = pd.DataFrame(
        {
            "polo": ["pimentas", "pimentas", "santana", "santana"],
            "team": ["Ana", "Ana", "Ana", "Ana"],
            "service": ["Água", "Esgoto", "Água", "Esgoto"],
            "conforme_count": [10, 8, 4, 12],
            "nao_conforme_count": [2, 4, 6, 0],
        }
    )

    rows_flat = indice_rows_from_inspections(df)
    rows_by_polo = indice_rows_from_inspections(df, group_by_polo=True)

    # Without the flag, Pimentas + Santana collapse into 2 (Ana, Água) and
    # (Ana, Esgoto) rows — the polo identity is lost.
    assert {r.equipe for r in rows_flat} == {"Ana"}
    # With the flag, each polo's Ana gets a distinct row.
    equipes = {r.equipe for r in rows_by_polo}
    assert equipes == {"Pimentas — Ana", "Santana — Ana"}


def test_render_semanal_embeds_chart_images(tmp_path):
    path = make_minimal_pimentas(tmp_path, with_inspections=True)
    wb = load_workbook(path, data_only=True, read_only=True)
    template = PimentasTemplate.detect(wb.sheetnames)

    ctx = semanal_context_from_template(template, path)
    body = render_semanal(ctx)

    # The skeleton already carries dozens of cover-page logos and body
    # screenshots; the assertion that matters is *delta*: three new image
    # relationships beyond the skeleton baseline, one per embedded chart.
    skeleton = Document(str(SEMANAL_SKELETON_PATH))
    baseline = sum("image" in r.reltype for r in skeleton.part.rels.values())

    doc = Document(BytesIO(body))
    rendered = sum("image" in r.reltype for r in doc.part.rels.values())
    assert rendered == baseline + 3

    text = _all_text(doc)
    for key in ("ic_bar", "iqs_bar", "photo_conformity"):
        assert "{{ " + key not in text
