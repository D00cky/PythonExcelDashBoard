"""Multi-file Polo report aggregation.

A *batch* is a directory holding 1..N weekly xlsx files (one Polo × one week per
file) plus a ``manifest.json`` sidecar that records the parsed period and Polo
for each file. This module owns the manifest format and the in-memory
representation; the route layer reads/writes batches through these primitives.
"""

from __future__ import annotations

import json
import re
from collections.abc import Iterable
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from app.core.templates import recognize
from app.core.templates.pimentas import PimentasTemplate, ServiceIC, ServiceIQS

MANIFEST_FILENAME = "manifest.json"

_PERIODO_RE = re.compile(
    r"(?P<d1>\d{2})/(?P<m1>\d{2})/(?P<y1>\d{4})\s*[àaÀ]\s*"
    r"(?P<d2>\d{2})/(?P<m2>\d{2})/(?P<y2>\d{4})"
)


def parse_periodo(s: str) -> tuple[date, date, str, str]:
    """Parse a ``"Período: dd/mm/yyyy à dd/mm/yyyy"`` cell into structured fields.

    Returns ``(start, end, iso_week, month)`` where ``iso_week`` is ``"YYYY-Www"``
    derived from the start date and ``month`` is ``"YYYY-MM"`` (also from start).
    """
    if not isinstance(s, str):
        raise TypeError(f"expected str, got {type(s).__name__}")
    match = _PERIODO_RE.search(s)
    if match is None:
        raise ValueError(f"no dd/mm/yyyy date range in {s!r}")
    start = date(int(match["y1"]), int(match["m1"]), int(match["d1"]))
    end = date(int(match["y2"]), int(match["m2"]), int(match["d2"]))
    iso_year, iso_week, _ = start.isocalendar()
    return start, end, f"{iso_year:04d}-W{iso_week:02d}", f"{start.year:04d}-{start.month:02d}"


@dataclass(frozen=True)
class PoloFile:
    file_path: Path
    polo: str
    iso_week: str
    month: str
    period_start: date
    period_end: date


@dataclass
class PoloBatch:
    batch_dir: Path
    files: list[PoloFile] = field(default_factory=list)

    @property
    def polos(self) -> list[str]:
        return sorted({f.polo for f in self.files})

    @property
    def iso_weeks(self) -> list[str]:
        return sorted({f.iso_week for f in self.files})

    @property
    def months(self) -> list[str]:
        return sorted({f.month for f in self.files})


def _period_from_inspections(
    template: PimentasTemplate, xlsx_path: Path
) -> tuple[date, date] | None:
    """Authoritative period from the inspections' ``start_date`` column.

    Returns ``None`` when no dated rows are present — the caller falls back to
    B4. We prefer this over B4 because users frequently clone last week's xlsx
    and forget to update the period cell, while the inspection-row dates are
    typed in fresh each cycle.
    """
    inspections = template.extract_inspections(xlsx_path)
    if "start_date" not in inspections.columns:
        return None
    dates = inspections["start_date"].dropna()
    if dates.empty:
        return None
    return dates.min().date(), dates.max().date()


def discover_polo_file(xlsx_path: Path) -> PoloFile:
    """Open an xlsx, detect its Polo template, and return a manifest entry.

    The period (start, end, iso_week, month) is derived from the inspections'
    ``start_date`` column when that column has any rows; otherwise we fall back
    to the B4 period cell. B4 lies often enough (cloned-from-last-week xlsx)
    that trusting it unconditionally surfaces as bug reports like "shows March
    for a May file".

    Raises ``ValueError`` if the workbook does not match any known template or
    if neither inspections nor B4 yield a parseable period.
    """
    workbook = load_workbook(xlsx_path, data_only=True, read_only=True)
    try:
        template = recognize(workbook.sheetnames)
        if template is None:
            raise ValueError(f"unknown template in {xlsx_path.name}")

        inspections_period = _period_from_inspections(template, xlsx_path)
        if inspections_period is not None:
            start, end = inspections_period
        else:
            raw = template.extract_periodo(workbook)
            if not raw:
                raise ValueError(f"no dated inspections and empty B4 in {xlsx_path.name}")
            start, end, _, _ = parse_periodo(raw)

        iso_year, iso_wk, _ = start.isocalendar()
        iso_week = f"{iso_year:04d}-W{iso_wk:02d}"
        month = f"{start.year:04d}-{start.month:02d}"
        return PoloFile(
            file_path=xlsx_path,
            polo=template.polo_name,
            iso_week=iso_week,
            month=month,
            period_start=start,
            period_end=end,
        )
    finally:
        workbook.close()


_VALID_VIEWS = ("weekly", "monthly")


def filter_batch(
    batch: PoloBatch, *, polos: tuple[str, ...], view: str, period_key: str
) -> PoloBatch:
    """Return a sub-batch matching ``polos`` and the ``view``/``period_key`` bucket."""
    if view not in _VALID_VIEWS:
        raise ValueError(f"view must be one of {_VALID_VIEWS}, got {view!r}")
    polos_set = set(polos)
    key_attr = "iso_week" if view == "weekly" else "month"
    selected = [
        f for f in batch.files if f.polo in polos_set and getattr(f, key_attr) == period_key
    ]
    return PoloBatch(batch_dir=batch.batch_dir, files=selected)


def combined_inspections(batch: PoloBatch) -> pd.DataFrame:
    """Concat per-file inspections with an extra ``polo`` column."""
    frames = []
    for f in batch.files:
        df = PimentasTemplate().extract_inspections(f.file_path)
        if df.empty:
            continue
        df = df.assign(polo=f.polo)
        frames.append(df)
    if not frames:
        return pd.DataFrame(columns=["polo"])
    return pd.concat(frames, ignore_index=True)


def iqs_rows_from_inspections(df: pd.DataFrame, services: list[str]) -> list[ServiceIQS]:
    """Reconstruct ServiceIQS records from raw inspection rows.

    Photos are summed across stage cells per service; NC + SF are lumped into
    ``fotos_nc`` so the result matches how DADOS aggregates failures.
    """
    out: list[ServiceIQS] = []
    if df.empty:
        return out
    for svc in services:
        sub = df[df["service"] == svc]
        if sub.empty:
            continue
        avaliadas = int(sub["photo_total"].sum())
        if avaliadas == 0:
            continue
        nc = int((sub["photo_nc"] + sub["photo_sf"]).sum())
        conforme = int(sub["photo_conforme"].sum())
        out.append(
            ServiceIQS(
                name=svc.title(),
                fotos_avaliadas=avaliadas,
                fotos_nc=nc,
                fotos_conforme=conforme,
                nc_pct=nc / avaliadas,
                conforme_pct=conforme / avaliadas,
            )
        )
    return out


def ic_rows_from_inspections(df: pd.DataFrame, services: list[str]) -> list[ServiceIC]:
    out: list[ServiceIC] = []
    if df.empty:
        return out
    for svc in services:
        sub = df[df["service"] == svc]
        total = len(sub)
        if total == 0:
            continue
        conf = int(sub["conforme_count"].sum())
        out.append(ServiceIC(name=svc.title(), ic_pct=conf / total, lvs=total))
    return out


def iqs_overall_from_inspections(df: pd.DataFrame) -> float | None:
    if df.empty:
        return None
    total = int(df["photo_total"].sum())
    if total == 0:
        return None
    return int(df["photo_conforme"].sum()) / total


def combined_iqs_rows(batch: PoloBatch, *, services: list[str]) -> list[ServiceIQS]:
    return iqs_rows_from_inspections(combined_inspections(batch), services)


def combined_ic_rows(batch: PoloBatch, *, services: list[str]) -> list[ServiceIC]:
    return ic_rows_from_inspections(combined_inspections(batch), services)


def combined_iqs_overall(batch: PoloBatch) -> float | None:
    return iqs_overall_from_inspections(combined_inspections(batch))


def combined_stage_failures(batch: PoloBatch) -> pd.DataFrame:
    """Concat per-file stage-failures with an extra ``polo`` column."""
    frames = []
    for f in batch.files:
        df = PimentasTemplate().extract_stage_failures(f.file_path)
        if df.empty:
            continue
        df = df.assign(polo=f.polo)
        frames.append(df)
    if not frames:
        return pd.DataFrame(columns=["polo"])
    return pd.concat(frames, ignore_index=True)


def write_manifest(batch_dir: Path, files: Iterable[PoloFile]) -> Path:
    payload = {
        "files": [
            {
                "file_name": f.file_path.name,
                "polo": f.polo,
                "iso_week": f.iso_week,
                "month": f.month,
                "period_start": f.period_start.isoformat(),
                "period_end": f.period_end.isoformat(),
            }
            for f in files
        ],
    }
    path = batch_dir / MANIFEST_FILENAME
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    return path


def load_batch(batch_dir: Path) -> PoloBatch:
    manifest_path = batch_dir / MANIFEST_FILENAME
    if not manifest_path.exists():
        raise FileNotFoundError(manifest_path)
    payload = json.loads(manifest_path.read_text(encoding="utf-8"))
    files = [
        PoloFile(
            file_path=batch_dir / entry["file_name"],
            polo=entry["polo"],
            iso_week=entry["iso_week"],
            month=entry["month"],
            period_start=date.fromisoformat(entry["period_start"]),
            period_end=date.fromisoformat(entry["period_end"]),
        )
        for entry in payload["files"]
    ]
    return PoloBatch(batch_dir=batch_dir, files=files)
