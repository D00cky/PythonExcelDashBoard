"""PowerPoint (.pptx) exporter — one slide per chart for projector use.

Builds slides from the same single-Polo data the dashboard renders. Charts
are emitted as PNGs via plotly+kaleido and inserted as full-slide images.
"""

from __future__ import annotations

import io
from pathlib import Path

from openpyxl.workbook import Workbook
from pptx import Presentation
from pptx.util import Inches, Pt

from app.core.aggregator import (
    ic_rows_from_inspections,
    iqs_overall_from_inspections,
    iqs_rows_from_inspections,
)
from app.core.templates.pimentas import PimentasTemplate


def render_pptx(template: PimentasTemplate, workbook: Workbook, path: Path) -> bytes:
    inspections = template.extract_inspections(path)
    failures = template.extract_stage_failures(path)
    services = sorted(template.SERVICE_SHEETS)
    ic_rows = template.extract_ic_by_service(workbook) or ic_rows_from_inspections(
        inspections, services
    )
    iqs_rows = template.extract_iqs_by_service(workbook) or iqs_rows_from_inspections(
        inspections, services
    )
    iqs_overall = template.extract_iqs_overall(workbook)
    if iqs_overall is None:
        iqs_overall = iqs_overall_from_inspections(inspections)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    _cover_slide(prs, template, workbook, iqs_overall, len(inspections))

    if ic_rows:
        _chart_slide(prs, "Índice de Conformidade", template.build_ic_bar(ic_rows))
    if iqs_rows:
        _chart_slide(prs, "Índice de Qualidade (IQS)", template.build_service_iqs_bar(iqs_rows))
        _chart_slide(
            prs,
            "Fotos por Serviço — Conforme vs NC",
            template.build_photo_conformity_stacked(iqs_rows),
        )
    if not inspections.empty:
        _chart_slide(
            prs, "Inspeções por Equipe × Serviço", template.build_team_service_stacked(inspections)
        )
        _chart_slide(prs, "TSS — Distribuição", template.build_tss_distribution(inspections))
        _chart_slide(prs, "Equipes com Mais NCs", template.build_worst_teams(inspections))
    if not failures.empty:
        _chart_slide(
            prs, "Etapas com Mais Não Conformidades", template.build_top_failing_stages(failures)
        )

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _cover_slide(prs, template, workbook, iqs_overall, total_inspections):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    left = Inches(0.6)
    top = Inches(0.6)
    width = Inches(12)
    height = Inches(1)
    title_box = slide.shapes.add_textbox(left, top, width, height)
    tf = title_box.text_frame
    tf.text = f"Dashboard — Polo {template.polo_name.title()}"
    tf.paragraphs[0].runs[0].font.size = Pt(40)
    tf.paragraphs[0].runs[0].font.bold = True

    periodo = template.extract_periodo(workbook)
    info_lines = []
    if periodo:
        info_lines.append(f"Período: {periodo}")
    if iqs_overall is not None:
        info_lines.append(f"IQS Geral: {iqs_overall:.1%}")
    info_lines.append(f"Inspeções avaliadas: {total_inspections}")

    info_box = slide.shapes.add_textbox(left, Inches(2.0), width, Inches(4))
    tf = info_box.text_frame
    for i, line in enumerate(info_lines):
        if i == 0:
            tf.text = line
            tf.paragraphs[0].runs[0].font.size = Pt(24)
        else:
            p = tf.add_paragraph()
            p.text = line
            p.runs[0].font.size = Pt(24)


def _chart_slide(prs, title: str, figure) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title_box = slide.shapes.add_textbox(Inches(0.4), Inches(0.2), Inches(12.5), Inches(0.7))
    tf = title_box.text_frame
    tf.text = title
    tf.paragraphs[0].runs[0].font.size = Pt(28)
    tf.paragraphs[0].runs[0].font.bold = True

    png = figure.to_image(format="png", width=1600, height=900, scale=1)
    slide.shapes.add_picture(io.BytesIO(png), Inches(0.4), Inches(1.0), Inches(12.5), Inches(6.2))
