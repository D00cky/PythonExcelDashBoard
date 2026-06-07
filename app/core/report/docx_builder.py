"""Build scoped Word reports with embedded chart images."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Inches

from app.core.report.charts import render_charts_for_scope
from app.core.report.data import load_report_data
from app.core.report.scope import ReportScope


def build_docx(
    uuid: str,
    period: str,
    scope: ReportScope,
    scope_name: str | None,
    output_path: Path,
) -> Path:
    """Build a scoped DOCX report and return ``output_path``."""
    data = load_report_data(uuid, scope=scope, scope_name=scope_name, period=period)
    charts = render_charts_for_scope(
        uuid,
        scope,
        scope_name,
        "light",
        output_path.parent / "charts" / "light",
        period=period,
    )

    doc = Document()
    doc.add_heading(_title(scope, data.label), level=0)
    doc.add_paragraph(f"Período: {data.period}")
    doc.add_paragraph(f"Escopo: {data.label}")
    doc.add_paragraph(f"Inspeções: {len(data.inspections)}")
    if not data.inspections.empty:
        doc.add_paragraph(f"Fotos avaliadas: {int(data.inspections['photo_total'].sum())}")

    doc.add_heading("Resumo", level=1)
    _summary_table(doc, data.inspections)

    for heading, key in (
        ("Índice de Conformidade", "chart_ic"),
        ("Índice de Qualidade", "chart_iqs"),
        ("Volume de inspeções", "chart_volume"),
    ):
        doc.add_heading(heading, level=1)
        doc.add_picture(str(charts[key]), width=Inches(6.5))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return output_path


def _summary_table(doc: Document, inspections) -> None:
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    headers = ("Grupo", "Inspeções", "Conforme", "Não Conforme")
    for idx, header in enumerate(headers):
        table.rows[0].cells[idx].text = header
    group_col = "polo" if "polo" in inspections.columns else "service"
    if inspections.empty:
        return
    grouped = inspections.groupby(group_col).agg(
        inspecoes=("conforme_count", "size"),
        conforme=("conforme_count", "sum"),
        nao_conforme=("nao_conforme_count", "sum"),
    )
    for label, row in grouped.iterrows():
        cells = table.add_row().cells
        cells[0].text = str(label)
        cells[1].text = str(int(row["inspecoes"]))
        cells[2].text = str(int(row["conforme"]))
        cells[3].text = str(int(row["nao_conforme"]))


def _title(scope: ReportScope, label: str) -> str:
    if scope == ReportScope.CITY:
        return "Relatório de Auditoria — São Paulo"
    if scope == ReportScope.ZONE:
        return f"Relatório — {label}"
    return f"Relatório — {label}"
