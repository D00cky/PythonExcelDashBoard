"""Build scoped PowerPoint reports."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

from app.core.report.charts import render_charts_for_scope
from app.core.report.data import load_report_data
from app.core.report.scope import ReportScope


def build_pptx(
    uuid: str,
    period: str,
    scope: ReportScope,
    scope_name: str | None,
    output_path: Path,
) -> Path:
    """Build a scoped PPTX report and return ``output_path``."""
    data = load_report_data(uuid, scope=scope, scope_name=scope_name, period=period)
    charts = render_charts_for_scope(
        uuid,
        scope,
        scope_name,
        "light",
        output_path.parent / "charts" / "light",
        period=period,
    )
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    _title_slide(prs, scope, data.label, data.period, len(data.inspections))
    for title, path in (
        ("Índice de Conformidade", charts["chart_ic"]),
        ("Índice de Qualidade", charts["chart_iqs"]),
        ("Volume de inspeções", charts["chart_volume"]),
    ):
        _chart_slide(prs, title, path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)
    return output_path


def _title_slide(
    prs: Presentation, scope: ReportScope, label: str, period: str, total: int
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.7), Inches(0.8), Inches(12), Inches(2))
    tf = box.text_frame
    tf.text = (
        "Relatório de Auditoria — São Paulo"
        if scope == ReportScope.CITY
        else f"Relatório — {label}"
    )
    tf.paragraphs[0].runs[0].font.size = Pt(36)
    tf.paragraphs[0].runs[0].font.bold = True
    p = tf.add_paragraph()
    p.text = f"Período: {period} · Inspeções: {total}"
    p.runs[0].font.size = Pt(20)


def _chart_slide(prs: Presentation, title: str, path: Path) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12), Inches(0.6))
    tf = box.text_frame
    tf.text = title
    tf.paragraphs[0].runs[0].font.size = Pt(28)
    tf.paragraphs[0].runs[0].font.bold = True
    slide.shapes.add_picture(str(path), Inches(0.7), Inches(1.1), Inches(11.9), Inches(5.8))
